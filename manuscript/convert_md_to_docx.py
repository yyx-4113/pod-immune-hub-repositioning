#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Minimal markdown -> .docx converter using python-docx (no pandoc needed).

Handles: #/##/###/#### headings, | tables | (with --- separator skipped),
'- ' bullet lists, and **bold** inline runs. Good enough for a manuscript draft.
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(paragraph, text):
    """Split on **bold** and add runs with bold flag."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        else:
            paragraph.add_run(part)


def convert(md_path, docx_path):
    doc = Document()
    # sane base style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = open(md_path, encoding="utf-8").read().split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
        elif stripped.startswith("|"):
            # collect a table block, skip the --- separator row
            tbl = []
            while i < n and lines[i].strip().startswith("|"):
                if re.match(r"^\s*\|[\s:\-\|]+\|\s*$", lines[i]):
                    i += 1
                    continue
                tbl.append(lines[i])
                i += 1
            if tbl:
                rows = []
                for r in tbl:
                    body = r.strip()
                    body = body[1:] if body.startswith("|") else body
                    body = body[:-1] if body.endswith("|") else body
                    # split on unescaped pipes only, so markdown-escaped \| inside
                    # a cell (e.g. \|t\| enrichment) does not create extra columns
                    rows.append([c.strip().replace("\\|", "|")
                                 for c in re.split(r"(?<!\\)\|", body)])
                if rows and rows[0]:
                    t = doc.add_table(rows=1, cols=len(rows[0]))
                    try:
                        t.style = "Light Grid Accent 1"
                    except Exception:
                        pass
                    hdr = t.rows[0].cells
                    for j, h in enumerate(rows[0]):
                        hdr[j].text = ""
                        add_runs(hdr[j].paragraphs[0], h)
                    for r in rows[1:]:
                        cells = t.add_row().cells
                        for j, c in enumerate(r):
                            cells[j].text = ""
                            add_runs(cells[j].paragraphs[0], c)
            continue
        m_img = re.match(r"^!\[(.*)\]\((.+)\)$", stripped)
        if m_img:
            alt, img = m_img.group(1), m_img.group(2).strip()
            img_path = img if os.path.isabs(img) else os.path.join(
                os.path.dirname(os.path.abspath(md_path)), img)
            p = doc.add_paragraph()
            try:
                p.add_run().add_picture(img_path, width=Inches(6.2))
            except Exception as e:
                p.add_run(f"[image missing: {img} ({e})]")
            if alt:
                cap = doc.add_paragraph()
                add_runs(cap, alt)
            i += 1
            continue
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:].strip())
        elif stripped == "":
            pass
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
        i += 1

    doc.save(docx_path)
    print("WROTE", docx_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
